#!/usr/bin/env python3
"""Convert UCI Online Retail II (xlsx, CC BY 4.0) to a slim CSV for bulk-load.mjs.
Usage: python3 prep_uci.py <in.xlsx> <out.csv> [sheet_index]   (sheet_index lets the two yearly sheets convert in parallel; src_row stays globally unique)
Keeps real product lines only (drops POST/DOT/BANK CHARGES-style service codes and price<=0)."""
import csv, re, sys
from openpyxl import load_workbook

src, dst = sys.argv[1], sys.argv[2]
only = int(sys.argv[3]) if len(sys.argv) > 3 else None
wb = load_workbook(src, read_only=True, data_only=True)
n = kept = 0
with open(dst, "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["src_row", "invoice", "stock_code", "description", "quantity", "invoice_ts", "price", "customer_id", "country"])
    for si, ws in enumerate(wb.worksheets):
        if only is not None and si != only:
            continue
        n = si * 2_000_000
        it = ws.iter_rows(values_only=True)
        next(it)  # header
        for r in it:
            n += 1
            inv, code, desc, qty, ts, price, cust, country = r[:8]
            if inv is None or code is None or ts is None or price is None or qty is None:
                continue
            code = str(code).strip().upper()
            if not re.match(r"^\d{4,6}", code):  # service codes: POST, DOT, M, BANK CHARGES, AMAZONFEE...
                continue
            if float(price) <= 0:
                continue
            w.writerow([n, str(inv), code, (desc or "").strip(), int(qty), ts.strftime("%Y-%m-%d %H:%M:%S"), f"{float(price):.4f}",
                        "" if cust is None else str(int(cust)), country or ""])
            kept += 1
            if n % 200000 == 0:
                print("rows", n, "kept", kept, flush=True)
print("done rows", n, "kept", kept)
